import streamlit as st
import anthropic
from datetime import datetime
import io
import os

# ── Identidade Visual ──────────────────────────────────────────────
NOME   = "Christiani Porto Gonçalves"
TITULO = "Nutricionista"
CRN    = "CRN-3: 75222"

COR_ROSA     = "#c97b7b"
COR_VERDE    = "#8bbcb0"
COR_BEGE     = "#f5ede8"
COR_BEGE_ESC = "#e8d5cc"
COR_TEXTO    = "#3a2e2a"

SENHA_APP = "nutri2025"  # ← altere aqui se quiser mudar a senha

REFEICOES = ["Café da manhã","Lanche da manhã","Almoço",
             "Lanche da tarde","Jantar","Ceia"]
HORARIOS_DEFAULT = {
    "Café da manhã":"07:00","Lanche da manhã":"10:00","Almoço":"12:00",
    "Lanche da tarde":"15:00","Jantar":"19:00","Ceia":"21:00",
}

# ── Página ─────────────────────────────────────────────────────────
st.set_page_config(
    page_title="NutriPlan — Christiani Porto Gonçalves",
    page_icon="🥗",
    layout="centered",
)

st.markdown(f"""
<style>
  @import url('https://fonts.googleapis.com/css2?family=Playfair+Display:wght@400;700&family=Lato:wght@300;400;700&display=swap');
  html, body, [class*="css"] {{ font-family: 'Lato', sans-serif; }}
  .header-box {{
    background: linear-gradient(135deg, {COR_BEGE} 0%, {COR_BEGE_ESC} 100%);
    border-bottom: 3px solid {COR_ROSA};
    padding: 20px 28px 14px;
    border-radius: 10px 10px 0 0;
    margin-bottom: 0;
  }}
  .header-nome {{
    font-family: 'Playfair Display', serif;
    font-size: 26px;
    font-weight: 700;
    color: {COR_TEXTO};
    margin: 0;
  }}
  .header-sub {{
    font-size: 11px;
    letter-spacing: 2px;
    color: {COR_VERDE};
    text-transform: uppercase;
    margin: 2px 0 0;
  }}
  .section-title {{
    font-size: 11px;
    letter-spacing: 1.5px;
    text-transform: uppercase;
    color: {COR_ROSA};
    font-weight: 700;
    margin: 24px 0 4px;
    border-bottom: 1px solid {COR_BEGE_ESC};
    padding-bottom: 4px;
  }}
  .resultado-box {{
    background: #fff;
    border: 1px solid {COR_BEGE_ESC};
    border-radius: 10px;
    padding: 24px;
    white-space: pre-wrap;
    font-size: 14px;
    line-height: 1.9;
    color: {COR_TEXTO};
  }}
  .calculo-box {{
    background: {COR_BEGE};
    border-left: 4px solid {COR_ROSA};
    border-radius: 6px;
    padding: 14px 18px;
    margin: 8px 0;
    font-size: 14px;
  }}
  .aviso {{
    font-size: 11px;
    color: #aaa;
    font-style: italic;
    margin-top: 10px;
  }}
  div[data-testid="stButton"] button {{
    background: {COR_ROSA};
    color: white;
    border: none;
    border-radius: 8px;
    padding: 10px 24px;
    font-size: 15px;
    font-weight: bold;
    font-family: 'Lato', sans-serif;
  }}
  div[data-testid="stButton"] button:hover {{
    background: #b06060;
  }}
</style>
""", unsafe_allow_html=True)

# ── Login ──────────────────────────────────────────────────────────
if "autenticado" not in st.session_state:
    st.session_state.autenticado = False

if not st.session_state.autenticado:
    st.markdown(f"""
    <div class="header-box">
      <p class="header-nome">{NOME}</p>
      <p class="header-sub">{TITULO} • {CRN}</p>
    </div>
    """, unsafe_allow_html=True)
    st.markdown("<br>", unsafe_allow_html=True)
    st.markdown("### 🔒 Acesso restrito")
    senha = st.text_input("Senha", type="password", placeholder="Digite a senha para acessar")
    if st.button("Entrar"):
        if senha == SENHA_APP:
            st.session_state.autenticado = True
            st.rerun()
        else:
            st.error("Senha incorreta.")
    st.stop()

# ── Header ─────────────────────────────────────────────────────────
st.markdown(f"""
<div class="header-box">
  <p class="header-nome">{NOME}</p>
  <p class="header-sub">{TITULO} • {CRN} &nbsp;|&nbsp; NutriPlan</p>
</div>
""", unsafe_allow_html=True)

st.markdown("<br>", unsafe_allow_html=True)

# ── Chave da API — Secrets (online) ou campo manual (local) ────────
try:
    api_key = st.secrets.get("ANTHROPIC_API_KEY", "")
except Exception:
    api_key = ""

with st.sidebar:
    st.markdown(f"### ⚙️ Configuração")
    if api_key:
        st.success("✅ Chave da API configurada")
    else:
        api_key = st.text_input(
            "Chave da API (Anthropic)",
            type="password",
            placeholder="sk-ant-...",
            help="Encontre em console.anthropic.com → API Keys")
    st.markdown("---")
    st.markdown(f"**{NOME}**")
    st.markdown(f"{TITULO} • {CRN}")
    st.markdown("---")
    if st.button("🚪 Sair"):
        st.session_state.autenticado = False
        st.rerun()

# ── Abas ───────────────────────────────────────────────────────────
aba1, aba2, aba3 = st.tabs(["👤 Paciente", "🎯 Metas e Refeições", "🥗 Plano Gerado"])

# ═══════════════════════════════════════════════
# ABA 1 — Paciente
# ═══════════════════════════════════════════════
with aba1:
    st.markdown('<p class="section-title">Dados Básicos</p>', unsafe_allow_html=True)

    col1, col2 = st.columns(2)
    with col1:
        sexo   = st.radio("Sexo", ["Feminino","Masculino"], horizontal=True)
        idade  = st.number_input("Idade (anos)", min_value=1, max_value=120, value=30)
    with col2:
        peso   = st.number_input("Peso (kg)", min_value=1.0, max_value=300.0, value=70.0, step=0.1)
        altura = st.number_input("Altura (cm)", min_value=50, max_value=250, value=165)

    # ── Cálculos automáticos ──
    st.markdown('<p class="section-title">Cálculos Automáticos</p>', unsafe_allow_html=True)

    # IMC
    altura_m = altura / 100
    imc = peso / (altura_m ** 2)
    if imc < 18.5:
        class_imc = "Abaixo do peso"
    elif imc < 25:
        class_imc = "Peso normal ✅"
    elif imc < 30:
        class_imc = "Sobrepeso"
    elif imc < 35:
        class_imc = "Obesidade grau I"
    elif imc < 40:
        class_imc = "Obesidade grau II"
    else:
        class_imc = "Obesidade grau III"

    # Peso ideal (Fórmula de Devine)
    if sexo == "Masculino":
        peso_ideal = 50 + 2.3 * ((altura - 152.4) / 2.54)
    else:
        peso_ideal = 45.5 + 2.3 * ((altura - 152.4) / 2.54)
    peso_ideal = max(peso_ideal, 30.0)

    # Gasto calórico basal — Mifflin-St Jeor
    if sexo == "Masculino":
        tmb = 10 * peso + 6.25 * altura - 5 * idade + 5
    else:
        tmb = 10 * peso + 6.25 * altura - 5 * idade - 161

    col_a, col_b, col_c = st.columns(3)
    with col_a:
        st.markdown(f"""
        <div class="calculo-box">
          <b>IMC</b><br>
          <span style="font-size:22px;color:{COR_ROSA};font-weight:700">{imc:.1f}</span><br>
          <span style="font-size:12px;color:#888">{class_imc}</span>
        </div>""", unsafe_allow_html=True)
    with col_b:
        st.markdown(f"""
        <div class="calculo-box">
          <b>Peso Ideal</b><br>
          <span style="font-size:22px;color:{COR_ROSA};font-weight:700">{peso_ideal:.1f} kg</span><br>
          <span style="font-size:12px;color:#888">Fórmula de Devine</span>
        </div>""", unsafe_allow_html=True)
    with col_c:
        st.markdown(f"""
        <div class="calculo-box">
          <b>Gasto Basal (TMB)</b><br>
          <span style="font-size:22px;color:{COR_ROSA};font-weight:700">{tmb:.0f} kcal</span><br>
          <span style="font-size:12px;color:#888">Mifflin-St Jeor</span>
        </div>""", unsafe_allow_html=True)

    # Gasto total por nível de atividade
    st.markdown('<p class="section-title">Gasto Calórico Total (GET)</p>', unsafe_allow_html=True)
    nivel = st.selectbox("Nível de atividade física", [
        "Sedentário (pouco ou nenhum exercício)",
        "Levemente ativo (exercício leve 1-3x/semana)",
        "Moderadamente ativo (exercício moderado 3-5x/semana)",
        "Muito ativo (exercício intenso 6-7x/semana)",
        "Extremamente ativo (exercício muito intenso, trabalho físico)",
    ])
    fatores = [1.2, 1.375, 1.55, 1.725, 1.9]
    fator = fatores[["Sedentário","Levemente","Moderadamente","Muito","Extremamente"]
                    .index(next(n for n in ["Sedentário","Levemente","Moderadamente","Muito","Extremamente"]
                                if n in nivel))]
    get = tmb * fator
    st.markdown(f"""
    <div class="calculo-box">
      Gasto Calórico Total Estimado: <b style="color:{COR_ROSA};font-size:18px">{get:.0f} kcal/dia</b>
      &nbsp;—&nbsp; TMB ({tmb:.0f}) × fator ({fator})
    </div>""", unsafe_allow_html=True)

    # ── Dieta atual ──
    st.markdown('<p class="section-title">Dieta Atual do Paciente</p>', unsafe_allow_html=True)
    st.caption("Descreva o que o paciente come atualmente (alimentos e quantidades por refeição)")
    dieta_atual = st.text_area("", height=180,
        placeholder="Café da manhã: 2 ovos mexidos, 2 fatias de pão integral, café com leite\n"
                    "Almoço: 2 conchas de arroz, 1 concha de feijão, frango grelhado 150g, salada\n"
                    "Jantar: macarrão com molho, suco de laranja",
        label_visibility="collapsed")

    # Salva na sessão
    st.session_state["dados"] = dict(
        sexo=sexo, idade=idade, peso=peso, altura=altura,
        imc=imc, class_imc=class_imc,
        peso_ideal=peso_ideal, tmb=tmb, get=get, nivel=nivel,
        dieta_atual=dieta_atual,
    )

# ═══════════════════════════════════════════════
# ABA 2 — Metas e Refeições
# ═══════════════════════════════════════════════
with aba2:
    st.markdown('<p class="section-title">Metas Nutricionais</p>', unsafe_allow_html=True)

    dados = st.session_state.get("dados", {})
    get_sugerido = int(dados.get("get", 2000))

    col1, col2 = st.columns(2)
    with col1:
        calorias  = st.number_input("Meta Calórica (kcal/dia)",
                                    min_value=500, max_value=5000,
                                    value=get_sugerido, step=50)
    with col2:
        proteinas = st.number_input("Meta de Proteínas (g/dia)",
                                    min_value=10, max_value=300, value=80, step=5)

    st.markdown('<p class="section-title">Refeições e Horários</p>', unsafe_allow_html=True)
    st.caption("Selecione as refeições e ajuste os horários")

    refeicoes_sel = []
    horarios_sel  = {}
    for r in REFEICOES:
        col_cb, col_hr = st.columns([3,1])
        with col_cb:
            sel = st.checkbox(r, key=f"ref_{r}")
        with col_hr:
            hr = st.text_input("", value=HORARIOS_DEFAULT[r],
                               key=f"hr_{r}", label_visibility="collapsed")
        if sel:
            refeicoes_sel.append(r)
            horarios_sel[r] = hr

    st.markdown('<p class="section-title">Restrições</p>', unsafe_allow_html=True)
    restricoes = st.text_input("O que não quer na dieta",
        placeholder="ex: não gosta de peixe, intolerante à lactose, vegetariano...",
        label_visibility="collapsed")

    st.session_state["metas"] = dict(
        calorias=calorias, proteinas=proteinas,
        refeicoes_sel=refeicoes_sel, horarios_sel=horarios_sel,
        restricoes=restricoes,
    )

# ═══════════════════════════════════════════════
# ABA 3 — Plano Gerado
# ═══════════════════════════════════════════════
with aba3:
    st.markdown('<p class="section-title">Gerar Plano Alimentar</p>', unsafe_allow_html=True)

    if st.button("🥗  Gerar Plano Alimentar"):
        dados  = st.session_state.get("dados",  {})
        metas  = st.session_state.get("metas",  {})

        if not api_key:
            st.error("Preencha a Chave da API na barra lateral.")
            st.stop()
        if not dados.get("dieta_atual","").strip():
            st.error("Descreva a dieta atual do paciente na aba 'Paciente'.")
            st.stop()
        if not metas.get("refeicoes_sel"):
            st.error("Selecione pelo menos uma refeição na aba 'Metas e Refeições'.")
            st.stop()

        refeicoes_list = ", ".join(
            f"{r} às {metas['horarios_sel'][r]}" for r in metas["refeicoes_sel"])

        prompt = f"""Você é um nutricionista especializado em planos alimentares baseados em evidências científicas brasileiras.

DADOS DO PACIENTE:
- Sexo: {dados['sexo']}
- Idade: {dados['idade']} anos
- Peso: {dados['peso']} kg
- Altura: {dados['altura']} cm
- IMC: {dados['imc']:.1f} ({dados['class_imc']})
- Peso ideal (Devine): {dados['peso_ideal']:.1f} kg
- TMB (Mifflin-St Jeor): {dados['tmb']:.0f} kcal
- Gasto calórico estimado: {dados['get']:.0f} kcal/dia ({dados['nivel']})

DIETA ATUAL DO PACIENTE:
{dados['dieta_atual']}

METAS NUTRICIONAIS:
- Calorias: {metas['calorias']} kcal/dia
- Proteínas: {metas['proteinas']} g/dia
- Refeições: {refeicoes_list}

RESTRIÇÕES: {metas['restricoes'] or 'Nenhuma'}

INSTRUÇÕES:
1. Elabore plano alimentar DIÁRIO baseado no que o paciente JÁ COME, adaptando quantidades.
2. Use apenas alimentos encontrados no Brasil.
3. Quantidades em medidas caseiras (colheres, conchas, xícaras, unidades, gramas).
4. Para cada alimento, informe calorias aproximadas (TACO/IBGE), proteínas e carboidratos.
5. Para cada alimento principal, sugira 1-2 substituições entre parênteses.
   Ex: Arroz branco (pode substituir por: arroz integral ou quinoa)
6. Total de calorias, proteínas e carboidratos por refeição e total diário.
7. Seção final com totais do dia e referências bibliográficas.

Formato de cada refeição:
## [NOME DA REFEIÇÃO] — [HORÁRIO]
- Alimento | quantidade | calorias kcal | proteínas g | carboidratos g
  (pode substituir por: opção 1 ou opção 2)
TOTAL: X kcal | Xg proteína | Xg carboidrato

## TOTAIS DO DIA
## REFERÊNCIAS BIBLIOGRÁFICAS"""

        with st.spinner("Gerando plano alimentar... aguarde (1-2 minutos)"):
            try:
                client = anthropic.Anthropic(api_key=api_key)
                msg = client.messages.create(
                    model="claude-opus-4-5",
                    max_tokens=4096,
                    messages=[{"role":"user","content":prompt}]
                )
                st.session_state["resultado"] = msg.content[0].text
                st.success("✅ Plano gerado com sucesso!")
            except Exception as e:
                st.error(f"Erro ao gerar plano: {e}")
                st.stop()

    # Exibe resultado
    resultado = st.session_state.get("resultado","")
    if resultado:
        st.markdown(resultado)
        st.markdown('<p class="aviso">⚠️ Plano gerado com auxílio de IA. Revise antes de prescrever.</p>',
                    unsafe_allow_html=True)

        col_w, col_x = st.columns(2)

        # ── Exportar Word ──
        with col_w:
            if st.button("📝 Exportar Word (.docx)"):
                from docx import Document
                from docx.shared import Pt, RGBColor, Cm
                from docx.enum.text import WD_ALIGN_PARAGRAPH
                from docx.oxml.ns import qn
                from docx.oxml import OxmlElement

                dados = st.session_state.get("dados", {})
                metas = st.session_state.get("metas", {})

                def hex_rgb(h):
                    h = h.lstrip("#")
                    return RGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))

                def set_cell_bg(cell, hex_color):
                    tc = cell._tc
                    tcPr = tc.get_or_add_tcPr()
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), hex_color.lstrip("#"))
                    tcPr.append(shd)

                def set_cell_margins(cell, top=80, bottom=80, left=150, right=150):
                    tc = cell._tc
                    tcPr = tc.get_or_add_tcPr()
                    tcMar = OxmlElement("w:tcMar")
                    for side, val in [("top",top),("bottom",bottom),("left",left),("right",right)]:
                        node = OxmlElement(f"w:{side}")
                        node.set(qn("w:w"), str(val))
                        node.set(qn("w:type"), "dxa")
                        tcMar.append(node)
                    tcPr.append(tcMar)

                def remove_borders(cell):
                    tc = cell._tc
                    tcPr = tc.get_or_add_tcPr()
                    tcB = OxmlElement("w:tcBorders")
                    for side in ["top","left","bottom","right","insideH","insideV"]:
                        n = OxmlElement(f"w:{side}")
                        n.set(qn("w:val"), "none")
                        tcB.append(n)
                    tcPr.append(tcB)

                doc = Document()
                for sec in doc.sections:
                    sec.top_margin = sec.bottom_margin = Cm(0)
                    sec.left_margin = sec.right_margin = Cm(0)

                # Cabeçalho
                from docx.shared import Inches
                t = doc.add_table(rows=1, cols=2)
                t.autofit = False
                t.columns[0].width = Inches(4.8)
                t.columns[1].width = Inches(1.8)
                for ci in range(2):
                    c = t.cell(0,ci)
                    set_cell_bg(c, COR_BEGE_ESC)
                    set_cell_margins(c, top=200, bottom=200,
                                     left=400 if ci==0 else 100,
                                     right=100 if ci==0 else 300)
                    remove_borders(c)

                p = t.cell(0,0).paragraphs[0]
                r = p.add_run(NOME)
                r.font.size=Pt(18); r.font.bold=True
                r.font.color.rgb=hex_rgb(COR_TEXTO); r.font.name="Georgia"
                p2 = t.cell(0,0).add_paragraph()
                r2 = p2.add_run(f"{TITULO}  •  {CRN}")
                r2.font.size=Pt(10); r2.font.color.rgb=hex_rgb(COR_VERDE); r2.font.name="Georgia"

                pd = t.cell(0,1).paragraphs[0]
                pd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                rd = pd.add_run(datetime.now().strftime("%d/%m/%Y"))
                rd.font.size=Pt(9); rd.font.color.rgb=hex_rgb("#aaaaaa"); rd.font.name="Georgia"

                # Linha rosa
                tl = doc.add_table(rows=1, cols=1)
                tl.autofit=False; tl.columns[0].width=Inches(6.6)
                cl = tl.cell(0,0)
                set_cell_bg(cl, COR_ROSA)
                set_cell_margins(cl, top=25, bottom=25, left=0, right=0)
                remove_borders(cl)
                cl.paragraphs[0].add_run("")

                # Faixa verde
                selecionadas = metas.get("refeicoes_sel",[])
                horarios_sel = metas.get("horarios_sel",{})
                ref_str = "  •  ".join(f"{r} às {horarios_sel.get(r,'')}" for r in selecionadas)
                tv = doc.add_table(rows=1, cols=2)
                tv.autofit=False
                tv.columns[0].width=Inches(3.3); tv.columns[1].width=Inches(3.3)
                for ci in range(2):
                    cv = tv.cell(0,ci)
                    set_cell_bg(cv, COR_VERDE)
                    set_cell_margins(cv, top=120, bottom=120, left=350, right=200)
                    remove_borders(cv)

                def verde_text(cell, label, value):
                    p = cell.paragraphs[0]
                    rb = p.add_run(label); rb.font.bold=True
                    rb.font.size=Pt(9); rb.font.color.rgb=RGBColor(255,255,255); rb.font.name="Georgia"
                    rv = p.add_run(value)
                    rv.font.size=Pt(9); rv.font.color.rgb=RGBColor(255,255,255); rv.font.name="Georgia"

                verde_text(tv.cell(0,0), "Paciente: ",
                    f"{dados.get('sexo','')}, {dados.get('idade','')} anos  •  "
                    f"{dados.get('peso','')} kg  •  {dados.get('altura','')} cm")
                verde_text(tv.cell(0,1), "Meta: ",
                    f"{metas.get('calorias','')} kcal/dia  •  {metas.get('proteinas','')} g proteína/dia")

                # Faixa bege
                tb = doc.add_table(rows=1, cols=1)
                tb.autofit=False; tb.columns[0].width=Inches(6.6)
                cb = tb.cell(0,0)
                set_cell_bg(cb, COR_BEGE)
                set_cell_margins(cb, top=80, bottom=80, left=350, right=200)
                remove_borders(cb)
                pb = cb.paragraphs[0]
                rbl = pb.add_run("Refeições: "); rbl.font.bold=True
                rbl.font.size=Pt(8); rbl.font.color.rgb=hex_rgb("#888888"); rbl.font.name="Georgia"
                rvl = pb.add_run(ref_str)
                rvl.font.size=Pt(8); rvl.font.color.rgb=hex_rgb("#888888"); rvl.font.name="Georgia"

                # Ajusta margens para conteúdo
                for sec in doc.sections:
                    sec.left_margin=Cm(2.5); sec.right_margin=Cm(2.5)

                doc.add_paragraph()

                # Conteúdo — reconhece tabelas markdown
                linhas = resultado.split("\n")
                i_linha = 0
                while i_linha < len(linhas):
                    s = linhas[i_linha].strip()

                    # Detecta bloco de tabela markdown
                    if s.startswith("|") and i_linha+1 < len(linhas) and linhas[i_linha+1].strip().startswith("|---"):
                        # Coleta todas as linhas da tabela
                        tbl_linhas = []
                        while i_linha < len(linhas) and linhas[i_linha].strip().startswith("|"):
                            row_s = linhas[i_linha].strip()
                            if not all(c in "-|: " for c in row_s):  # ignora separador
                                cols = [c.strip().strip("*") for c in row_s.split("|") if c.strip()]
                                tbl_linhas.append(cols)
                            i_linha += 1

                        if tbl_linhas:
                            max_cols = max(len(r) for r in tbl_linhas)
                            from docx.shared import Inches
                            tbl = doc.add_table(rows=len(tbl_linhas), cols=max_cols)
                            tbl.style = "Table Grid"
                            from docx.oxml.ns import qn as qn2
                            from docx.oxml import OxmlElement as OE2
                            for ri, row_data in enumerate(tbl_linhas):
                                for ci, val in enumerate(row_data):
                                    if ci >= max_cols: break
                                    cell = tbl.cell(ri, ci)
                                    # Cor de fundo alternada
                                    tc = cell._tc
                                    tcPr = tc.get_or_add_tcPr()
                                    shd = OE2("w:shd")
                                    shd.set(qn2("w:val"), "clear")
                                    shd.set(qn2("w:color"), "auto")
                                    bg = "E8D5CC" if ri == 0 else ("FFFFFF" if ri%2==0 else "F5EDE8")
                                    shd.set(qn2("w:fill"), bg)
                                    tcPr.append(shd)
                                    # Texto
                                    is_sub = "(pode substituir" in val.lower()
                                    p = cell.paragraphs[0]
                                    run = p.add_run(val)
                                    run.font.name = "Georgia"
                                    run.font.size = Pt(8)
                                    run.font.bold = (ri == 0)
                                    run.font.color.rgb = hex_rgb(COR_VERDE if is_sub else COR_TEXTO)
                                    run.font.italic = is_sub
                            doc.add_paragraph()
                        continue

                    if not s:
                        i_linha += 1
                        continue
                    if s.startswith("## "):
                        doc.add_paragraph()
                        p = doc.add_paragraph()
                        r = p.add_run(s[3:].strip())
                        r.font.size=Pt(12); r.font.bold=True
                        r.font.color.rgb=hex_rgb(COR_ROSA); r.font.name="Georgia"
                        sep = doc.add_paragraph()
                        sr = sep.add_run("─"*55)
                        sr.font.size=Pt(7); sr.font.color.rgb=hex_rgb(COR_BEGE_ESC)
                    elif s.startswith("# "):
                        p = doc.add_paragraph()
                        r = p.add_run(s[2:].strip())
                        r.font.size=Pt(14); r.font.bold=True
                        r.font.color.rgb=hex_rgb(COR_ROSA); r.font.name="Georgia"
                    elif s.startswith("- ") or s.startswith("* "):
                        texto = s[2:]
                        p = doc.add_paragraph()
                        p.paragraph_format.left_indent=Cm(0.5)
                        p.paragraph_format.space_after=Pt(1)
                        if "(pode substituir" in texto.lower():
                            partes = texto.split("(pode substituir")
                            rb = p.add_run("• " + partes[0].strip() + " ")
                            rb.font.size=Pt(10); rb.font.name="Georgia"
                            rs = p.add_run("(pode substituir" + partes[1])
                            rs.font.size=Pt(9); rs.font.italic=True
                            rs.font.color.rgb=hex_rgb(COR_VERDE); rs.font.name="Georgia"
                        else:
                            rb = p.add_run("• " + texto)
                            rb.font.size=Pt(10); rb.font.name="Georgia"
                    else:
                        p = doc.add_paragraph()
                        p.paragraph_format.space_after=Pt(2)
                        partes = s.split("**")
                        for j, parte in enumerate(partes):
                            if not parte: continue
                            r = p.add_run(parte)
                            r.font.size=Pt(10); r.font.bold=(j%2==1); r.font.name="Georgia"
                    i_linha += 1

                # Assinatura
                doc.add_paragraph()
                sep2 = doc.add_paragraph("─"*70)
                sep2.runs[0].font.size=Pt(7)
                sep2.runs[0].font.color.rgb=hex_rgb(COR_BEGE_ESC)
                av = doc.add_paragraph()
                rav = av.add_run("Plano elaborado com auxílio de IA. Revise antes de prescrever.")
                rav.font.size=Pt(8); rav.font.italic=True
                rav.font.color.rgb=hex_rgb("#aaaaaa"); rav.font.name="Georgia"
                doc.add_paragraph()
                assin = doc.add_paragraph()
                assin.alignment=WD_ALIGN_PARAGRAPH.RIGHT
                assin.paragraph_format.space_before=Pt(20)
                assin.add_run("_"*32+"\n").font.size=Pt(10)
                rn = assin.add_run(NOME+"\n"); rn.bold=True
                rn.font.size=Pt(10); rn.font.name="Georgia"
                rc = assin.add_run(f"{TITULO}  •  {CRN}")
                rc.font.size=Pt(9); rc.font.color.rgb=hex_rgb("#aaaaaa"); rc.font.name="Georgia"

                buf = io.BytesIO()
                doc.save(buf)
                buf.seek(0)
                st.download_button(
                    "⬇️ Baixar Word",
                    data=buf,
                    file_name=f"PlanoAlimentar_{datetime.now().strftime('%Y%m%d_%H%M')}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )

        # ── Exportar Excel ──
        with col_x:
            if st.button("📊 Exportar Excel (.xlsx)"):
                import openpyxl
                from openpyxl.styles import (Font, PatternFill, Alignment,
                                              Border, Side, numbers)
                from openpyxl.utils import get_column_letter

                dados = st.session_state.get("dados", {})
                metas = st.session_state.get("metas", {})

                wb = openpyxl.Workbook()

                # ── Aba 1: Plano Alimentar ──
                ws = wb.active
                ws.title = "Plano Alimentar"
                ws.sheet_view.showGridLines = False

                # Cores
                fill_rosa    = PatternFill("solid", fgColor="C97B7B")
                fill_bege    = PatternFill("solid", fgColor="F5EDE8")
                fill_bege_e  = PatternFill("solid", fgColor="E8D5CC")
                fill_verde   = PatternFill("solid", fgColor="8BBCB0")
                fill_branco  = PatternFill("solid", fgColor="FFFFFF")
                fill_sub     = PatternFill("solid", fgColor="EEF7F5")

                thin = Side(style="thin", color="E8D5CC")
                borda = Border(bottom=Side(style="thin", color="E8D5CC"))

                def hdr(ws, row, col, val, size=10, bold=False, fill=None,
                        color="3A2E2A", align="left", wrap=False):
                    c = ws.cell(row=row, column=col, value=val)
                    c.font = Font(name="Georgia", size=size, bold=bold, color=color)
                    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
                    if fill: c.fill = fill
                    return c

                # Cabeçalho
                ws.merge_cells("A1:F1")
                c = ws["A1"]
                c.value = NOME
                c.font = Font(name="Georgia", size=18, bold=True, color="3A2E2A")
                c.fill = fill_bege_e
                c.alignment = Alignment(horizontal="left", vertical="center")
                ws.row_dimensions[1].height = 32

                ws.merge_cells("A2:F2")
                c2 = ws["A2"]
                c2.value = f"{TITULO}  •  {CRN}  |  Data: {datetime.now().strftime('%d/%m/%Y')}"
                c2.font = Font(name="Georgia", size=10, color="8BBCB0")
                c2.fill = fill_bege_e
                c2.alignment = Alignment(horizontal="left", vertical="center")
                ws.row_dimensions[2].height = 18

                # Linha rosa
                ws.merge_cells("A3:F3")
                ws["A3"].fill = fill_rosa
                ws.row_dimensions[3].height = 4

                # Dados do paciente
                ws.merge_cells("A4:F4")
                ws["A4"].value = (
                    f"Paciente: {dados.get('sexo','')}, {dados.get('idade','')} anos  •  "
                    f"{dados.get('peso','')} kg  •  {dados.get('altura','')} cm  •  "
                    f"IMC: {dados.get('imc',0):.1f} ({dados.get('class_imc','')})  •  "
                    f"Peso ideal: {dados.get('peso_ideal',0):.1f} kg  •  "
                    f"TMB: {dados.get('tmb',0):.0f} kcal  •  GET: {dados.get('get',0):.0f} kcal/dia"
                )
                ws["A4"].font = Font(name="Georgia", size=9, color="FFFFFF", bold=True)
                ws["A4"].fill = fill_verde
                ws["A4"].alignment = Alignment(horizontal="left", vertical="center")
                ws.row_dimensions[4].height = 20

                ws.merge_cells("A5:F5")
                ws["A5"].value = (
                    f"Meta: {metas.get('calorias','')} kcal/dia  •  "
                    f"{metas.get('proteinas','')} g proteína/dia  •  "
                    f"Restrições: {metas.get('restricoes','') or 'Nenhuma'}"
                )
                ws["A5"].font = Font(name="Georgia", size=9, color="888888")
                ws["A5"].fill = fill_bege
                ws["A5"].alignment = Alignment(horizontal="left", vertical="center")
                ws.row_dimensions[5].height = 16

                ws.row_dimensions[6].height = 8

                # Cabeçalho da tabela
                headers = ["Refeição / Alimento","Quantidade","Calorias (kcal)","Proteínas (g)","Carboidratos (g)","Substituições"]
                for ci, h in enumerate(headers, 1):
                    c = ws.cell(row=7, column=ci, value=h)
                    c.font = Font(name="Georgia", size=10, bold=True, color="FFFFFF")
                    c.fill = fill_rosa
                    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    c.border = Border(bottom=Side(style="medium", color="B06060"))
                ws.row_dimensions[7].height = 28

                # Larguras
                ws.column_dimensions["A"].width = 32
                ws.column_dimensions["B"].width = 18
                ws.column_dimensions["C"].width = 16
                ws.column_dimensions["D"].width = 16
                ws.column_dimensions["E"].width = 18
                ws.column_dimensions["F"].width = 36

                # Parseia o resultado — reconhece tabelas markdown
                import re as _re
                def extrair_num(txt):
                    m = _re.search(r"[\d]+[.,]?[\d]*", txt.replace(",","."))
                    return float(m.group().replace(",",".")) if m else ""

                row = 8
                linhas_res = resultado.split("\n")
                ir = 0
                refeicao_atual = None
                while ir < len(linhas_res):
                    s = linhas_res[ir].strip()

                    # Tabela markdown — cabeçalho seguido de separador |---|
                    if (s.startswith("|") and
                        ir+1 < len(linhas_res) and
                        linhas_res[ir+1].strip().startswith("|---")):

                        # Pula cabeçalho e separador
                        ir += 2
                        # Lê linhas de dados da tabela
                        while ir < len(linhas_res) and linhas_res[ir].strip().startswith("|"):
                            row_s = linhas_res[ir].strip()
                            cols = [c.strip().strip("*") for c in row_s.split("|") if c.strip()]
                            if len(cols) >= 2:
                                alimento = cols[0] if len(cols) > 0 else ""
                                qtd      = cols[1] if len(cols) > 1 else ""
                                cal      = extrair_num(cols[2]) if len(cols) > 2 else ""
                                prot     = extrair_num(cols[3]) if len(cols) > 3 else ""
                                carb     = extrair_num(cols[4]) if len(cols) > 4 else ""
                                sub      = cols[5] if len(cols) > 5 else ""

                                # Ignora linhas só com "pode substituir"
                                if "(pode substituir" in alimento.lower():
                                    # coloca sub na linha anterior
                                    if row > 8:
                                        ws.cell(row=row-1, column=6).value = alimento
                                    ir += 1
                                    continue

                                fill_linha = fill_branco if row % 2 == 0 else fill_bege
                                # Arredonda números
                                def fmt_num(v):
                                    if isinstance(v, float):
                                        return round(v, 1)
                                    return v

                                for ci, val in enumerate([alimento, qtd,
                                                          fmt_num(cal),
                                                          fmt_num(prot),
                                                          fmt_num(carb),
                                                          sub], 1):
                                    c = ws.cell(row=row, column=ci, value=val)
                                    c.font = Font(name="Georgia", size=9,
                                                  italic=(ci==6),
                                                  color="8BBCB0" if ci==6 else "3A2E2A")
                                    c.fill = fill_sub if ci==6 else fill_linha
                                    c.alignment = Alignment(
                                        horizontal="center" if ci>1 else "left",
                                        vertical="center", wrap_text=True)
                                    c.border = Border(bottom=Side(style="thin", color="E8D5CC"))
                                    # Formato numérico para colunas de cálculo
                                    if ci in (3, 4, 5) and isinstance(fmt_num(val), float):
                                        c.number_format = "0.0"
                                ws.row_dimensions[row].height = 18
                                row += 1
                            ir += 1
                        continue

                    if not s:
                        ir += 1
                        continue

                    if s.startswith("## ") or s.startswith("# "):
                        titulo = s.lstrip("#").strip()
                        ws.merge_cells(f"A{row}:F{row}")
                        c = ws.cell(row=row, column=1, value=titulo)
                        c.font = Font(name="Georgia", size=11, bold=True, color="C97B7B")
                        c.fill = fill_bege_e
                        c.alignment = Alignment(horizontal="left", vertical="center")
                        c.border = Border(bottom=Side(style="thin", color="C97B7B"))
                        ws.row_dimensions[row].height = 22
                        refeicao_atual = titulo
                        row += 1

                    elif "TOTAL" in s.upper() and (":" in s or "|" in s):
                        clean = s.replace("**","").replace("🔸","").strip()
                        ws.merge_cells(f"A{row}:F{row}")
                        c = ws.cell(row=row, column=1, value=clean)
                        c.font = Font(name="Georgia", size=9, bold=True, color="3A2E2A")
                        c.fill = fill_bege_e
                        c.alignment = Alignment(horizontal="right", vertical="center")
                        c.border = Border(top=Side(style="thin", color="C97B7B"))
                        ws.row_dimensions[row].height = 16
                        row += 1

                    ir += 1

                # ── Aba 2: Cálculos ──
                ws2 = wb.create_sheet("Cálculos Nutricionais")
                ws2.sheet_view.showGridLines = False
                ws2.column_dimensions["A"].width = 30
                ws2.column_dimensions["B"].width = 22
                ws2.column_dimensions["C"].width = 30

                ws2.merge_cells("A1:C1")
                ws2["A1"].value = "Cálculos Nutricionais"
                ws2["A1"].font = Font(name="Georgia", size=14, bold=True, color="C97B7B")
                ws2["A1"].fill = fill_bege_e
                ws2["A1"].alignment = Alignment(horizontal="left", vertical="center")
                ws2.row_dimensions[1].height = 28

                calculos = [
                    ("DADOS DO PACIENTE", "", ""),
                    ("Sexo", dados.get("sexo",""), ""),
                    ("Idade", f"{dados.get('idade','')} anos", ""),
                    ("Peso atual", f"{dados.get('peso','')} kg", ""),
                    ("Altura", f"{dados.get('altura','')} cm", ""),
                    ("", "", ""),
                    ("RESULTADOS", "", ""),
                    ("IMC", f"{dados.get('imc',0):.1f}", dados.get("class_imc","")),
                    ("Peso Ideal", f"{dados.get('peso_ideal',0):.1f} kg", "Fórmula de Devine"),
                    ("TMB (Basal)", f"{dados.get('tmb',0):.0f} kcal/dia", "Mifflin-St Jeor"),
                    ("Nível de atividade", dados.get("nivel",""), ""),
                    ("GET (Gasto Total)", f"{dados.get('get',0):.0f} kcal/dia", "TMB × fator atividade"),
                    ("", "", ""),
                    ("METAS", "", ""),
                    ("Meta calórica", f"{metas.get('calorias','')} kcal/dia", ""),
                    ("Meta proteica", f"{metas.get('proteinas','')} g/dia", ""),
                    ("Diferença (meta vs GET)",
                     f"{metas.get('calorias',0) - int(dados.get('get',0) or 0):+.0f} kcal",
                     "Negativo = déficit | Positivo = superávit"),
                ]

                for i, (label, valor, obs) in enumerate(calculos, 2):
                    is_header = label in ["DADOS DO PACIENTE","RESULTADOS","METAS"]
                    ws2.cell(row=i, column=1, value=label).font = Font(
                        name="Georgia", size=10,
                        bold=is_header, color="C97B7B" if is_header else "3A2E2A")
                    ws2.cell(row=i, column=1).fill = fill_bege_e if is_header else fill_branco
                    ws2.cell(row=i, column=2, value=valor).font = Font(
                        name="Georgia", size=10, bold=not is_header, color="3A2E2A")
                    ws2.cell(row=i, column=2).fill = fill_branco
                    ws2.cell(row=i, column=3, value=obs).font = Font(
                        name="Georgia", size=9, italic=True, color="888888")
                    ws2.cell(row=i, column=3).fill = fill_branco
                    ws2.row_dimensions[i].height = 18

                buf2 = io.BytesIO()
                wb.save(buf2)
                buf2.seek(0)
                st.download_button(
                    "⬇️ Baixar Excel",
                    data=buf2,
                    file_name=f"PlanoAlimentar_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
